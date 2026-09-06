"""Import initial des données depuis le fichier Excel de départ.

Usage :
    python scripts/import_xlsx.py chemin/vers/fichier.xlsx [--dry-run]

Le mapping exact des colonnes Excel n'est pas encore connu : les constantes et les
fonctions `_row_to_*` ci-dessous sont marquées TODO et doivent être complétées une fois
le format du fichier fourni. Le reste (connexion DB, idempotence, logging) est en place.
"""

import argparse
import asyncio
import logging
import sys
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.db import async_session_maker  # noqa: E402
from app.models.category import Category  # noqa: E402
from app.models.loan import Loan, LoanRepayment  # noqa: E402
from app.models.payment_plan import PaymentPlan  # noqa: E402
from app.models.recurring_charge import RecurringCharge  # noqa: E402
from app.models.transaction import Transaction  # noqa: E402

logging.basicConfig(level=logging.INFO, format="%(levelname)-7s %(message)s")
logger = logging.getLogger("import_xlsx")

# --- Mapping du fichier Excel -------------------------------------------------
# TODO : renseigner les noms exacts des onglets une fois le fichier fourni.
SHEET_TRANSACTIONS = "TODO"
SHEET_PAYMENT_PLANS = "TODO"
SHEET_RECURRING_CHARGES = "TODO"
SHEET_LOANS = "TODO"

# TODO : mapping "libellé de catégorie dans le Excel" -> "nom de catégorie en base".
# Les catégories cibles sont celles créées par le seed (app/seed.py).
CATEGORY_LABEL_MAP: dict[str, str] = {}

# Catégorie de repli quand le libellé Excel n'est pas reconnu.
FALLBACK_EXPENSE_CATEGORY = "Autres"
FALLBACK_INCOME_CATEGORY = "Autres revenus"


@dataclass
class ImportStats:
    created: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)

    def log(self, label: str) -> None:
        logger.info("%s : %d créé(s), %d ignoré(s)", label, self.created, self.skipped)
        for error in self.errors:
            logger.warning("%s : %s", label, error)


class CategoryResolver:
    """Résout un libellé de catégorie Excel vers une catégorie en base."""

    def __init__(self, categories: list[Category]) -> None:
        self._by_name = {category.name: category for category in categories}

    def resolve(self, excel_label: str | None, kind: str) -> Category:
        mapped_name = CATEGORY_LABEL_MAP.get((excel_label or "").strip())
        if mapped_name and mapped_name in self._by_name:
            return self._by_name[mapped_name]

        fallback = FALLBACK_EXPENSE_CATEGORY if kind == "expense" else FALLBACK_INCOME_CATEGORY
        return self._by_name[fallback]


# --- Conversion ligne Excel -> modèle ----------------------------------------


def _row_to_transaction(row: tuple, resolver: CategoryResolver) -> Transaction | None:
    """TODO : construire une Transaction depuis une ligne de l'onglet dépenses/revenus.

    À déterminer une fois le fichier fourni : index/nom des colonnes, format des dates,
    façon dont le Excel distingue une dépense d'un revenu, et le statut (réglé ou non).
    """
    raise NotImplementedError("Mapping des colonnes à compléter")


def _row_to_payment_plan(row: tuple) -> PaymentPlan | None:
    """TODO : construire un PaymentPlan (créancier, total, nb de mensualités, mois de début)."""
    raise NotImplementedError("Mapping des colonnes à compléter")


def _row_to_recurring_charge(row: tuple, resolver: CategoryResolver) -> RecurringCharge | None:
    """TODO : construire une RecurringCharge (libellé, montant, catégorie, période d'activité)."""
    raise NotImplementedError("Mapping des colonnes à compléter")


def _row_to_loan_repayment(row: tuple, loan: Loan) -> LoanRepayment | None:
    """TODO : construire un LoanRepayment (date de versement, montant, note)."""
    raise NotImplementedError("Mapping des colonnes à compléter")


# --- Idempotence --------------------------------------------------------------


async def _transaction_exists(
    db: AsyncSession, kind: str, description: str, amount: Decimal, due_date: date | None
) -> bool:
    result = await db.execute(
        select(Transaction.id).where(
            Transaction.kind == kind,
            Transaction.description == description,
            Transaction.amount == amount,
            Transaction.due_date == due_date,
        )
    )
    return result.first() is not None


async def _payment_plan_exists(
    db: AsyncSession, creditor: str, total_amount: Decimal, start_month: date
) -> bool:
    result = await db.execute(
        select(PaymentPlan.id).where(
            PaymentPlan.creditor == creditor,
            PaymentPlan.total_amount == total_amount,
            PaymentPlan.start_month == start_month,
        )
    )
    return result.first() is not None


async def _recurring_charge_exists(
    db: AsyncSession, label: str, amount: Decimal, active_from: date
) -> bool:
    result = await db.execute(
        select(RecurringCharge.id).where(
            RecurringCharge.label == label,
            RecurringCharge.amount == amount,
            RecurringCharge.active_from == active_from,
        )
    )
    return result.first() is not None


async def _get_or_create_loan(db: AsyncSession, label: str, total_amount: Decimal) -> Loan:
    result = await db.execute(select(Loan).where(Loan.label == label))
    loan = result.scalar_one_or_none()
    if loan is not None:
        return loan

    loan = Loan(label=label, total_amount=total_amount)
    db.add(loan)
    await db.flush()
    logger.info("Prêt créé : %s", label)
    return loan


async def _loan_repayment_exists(
    db: AsyncSession, loan: Loan, payment_date: date, amount: Decimal
) -> bool:
    result = await db.execute(
        select(LoanRepayment.id).where(
            LoanRepayment.loan_id == loan.id,
            LoanRepayment.payment_date == payment_date,
            LoanRepayment.amount == amount,
        )
    )
    return result.first() is not None


# --- Import par onglet --------------------------------------------------------


def _data_rows(sheet: Worksheet) -> list[tuple]:
    """Lignes de données, en-tête exclu. TODO : ajuster si le Excel a plusieurs lignes d'en-tête."""
    return list(sheet.iter_rows(min_row=2, values_only=True))


async def import_transactions(
    db: AsyncSession, sheet: Worksheet, resolver: CategoryResolver
) -> ImportStats:
    stats = ImportStats()
    for index, row in enumerate(_data_rows(sheet), start=2):
        try:
            transaction = _row_to_transaction(row, resolver)
        except NotImplementedError:
            raise
        except Exception as exc:
            stats.errors.append(f"ligne {index} : {exc}")
            continue

        if transaction is None:
            stats.skipped += 1
            continue

        if await _transaction_exists(
            db,
            transaction.kind,
            transaction.description,
            transaction.amount,
            transaction.due_date,
        ):
            stats.skipped += 1
            continue

        db.add(transaction)
        stats.created += 1

    return stats


async def import_payment_plans(db: AsyncSession, sheet: Worksheet) -> ImportStats:
    stats = ImportStats()
    for index, row in enumerate(_data_rows(sheet), start=2):
        try:
            plan = _row_to_payment_plan(row)
        except NotImplementedError:
            raise
        except Exception as exc:
            stats.errors.append(f"ligne {index} : {exc}")
            continue

        if plan is None:
            stats.skipped += 1
            continue

        if await _payment_plan_exists(db, plan.creditor, plan.total_amount, plan.start_month):
            stats.skipped += 1
            continue

        db.add(plan)
        stats.created += 1

    return stats


async def import_recurring_charges(
    db: AsyncSession, sheet: Worksheet, resolver: CategoryResolver
) -> ImportStats:
    stats = ImportStats()
    for index, row in enumerate(_data_rows(sheet), start=2):
        try:
            charge = _row_to_recurring_charge(row, resolver)
        except NotImplementedError:
            raise
        except Exception as exc:
            stats.errors.append(f"ligne {index} : {exc}")
            continue

        if charge is None:
            stats.skipped += 1
            continue

        if await _recurring_charge_exists(db, charge.label, charge.amount, charge.active_from):
            stats.skipped += 1
            continue

        db.add(charge)
        stats.created += 1

    return stats


async def import_loans(db: AsyncSession, sheet: Worksheet) -> ImportStats:
    """TODO : le Excel décrit un prêt et ses versements — à confirmer une fois le fichier fourni.

    Structure attendue : un libellé + un montant total pour le prêt, puis une ligne par
    remboursement. `_get_or_create_loan` garantit qu'un re-run ne duplique pas le prêt.
    """
    stats = ImportStats()
    raise NotImplementedError("Structure de l'onglet prêt à compléter")


# --- Entrée -------------------------------------------------------------------


async def run_import(xlsx_path: Path, dry_run: bool) -> None:
    workbook = load_workbook(xlsx_path, data_only=True)
    logger.info("Fichier chargé : %s (onglets : %s)", xlsx_path.name, workbook.sheetnames)

    async with async_session_maker() as db:
        categories = (await db.execute(select(Category))).scalars().all()
        if not categories:
            logger.error("Aucune catégorie en base — lancer le backend une fois pour le seed.")
            return
        resolver = CategoryResolver(categories)

        sheets = {
            "Transactions": (SHEET_TRANSACTIONS, import_transactions, (resolver,)),
            "Plans de paiement": (SHEET_PAYMENT_PLANS, import_payment_plans, ()),
            "Charges récurrentes": (SHEET_RECURRING_CHARGES, import_recurring_charges, (resolver,)),
            "Prêts": (SHEET_LOANS, import_loans, ()),
        }

        for label, (sheet_name, importer, extra_args) in sheets.items():
            if sheet_name not in workbook.sheetnames:
                logger.warning("%s : onglet '%s' absent du fichier, ignoré.", label, sheet_name)
                continue

            stats = await importer(db, workbook[sheet_name], *extra_args)
            stats.log(label)

        if dry_run:
            await db.rollback()
            logger.info("Dry-run : aucune écriture en base.")
        else:
            await db.commit()
            logger.info("Import terminé.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Importe les données de démarrage depuis un .xlsx")
    parser.add_argument("xlsx_path", type=Path, help="Chemin vers le fichier Excel")
    parser.add_argument(
        "--dry-run", action="store_true", help="Simule l'import sans écrire en base"
    )
    args = parser.parse_args()

    if not args.xlsx_path.is_file():
        parser.error(f"Fichier introuvable : {args.xlsx_path}")

    asyncio.run(run_import(args.xlsx_path, args.dry_run))


if __name__ == "__main__":
    main()
